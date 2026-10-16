"""Compara, cifra por cifra, los dos caminos del oraculo de 5.2.

  Camino 1: oraculo_5_2.xlsx, hoja «Esperados» (formulas de Excel, doble precision),
            ya recalculado con recalc_excel.ps1.
  Camino 2: esperados/*.csv, rechazos_rf28.csv y casos_positivos.csv
            (calcula_oraculo.py, decimal de 28 digitos).

Empareja los renglones por llave (no por posicion) y reporta, por tipo de cifra,
la diferencia absoluta maxima y donde ocurre. Las fechas, textos, enteros y P de
CETES deben coincidir exactamente; las demas cifras, a 1e-6 o menos (se espera
1e-9 o menos, por la doble precision de la hoja). En rechazos_rf28 y
casos_positivos solo se comparan las columnas calculadas (las que tiene la hoja).
Escribe comparacion_vias.csv.

Uso: python -B <ruta del oraculo>/compara_vias.py   (sin -I: necesita openpyxl)
"""
import csv
import sys
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

RAIZ = Path(__file__).resolve().parent
LIBRO = RAIZ / "oraculo_5_2.xlsx"
ESP = RAIZ / "esperados"
UMBRAL = Decimal("1e-6")

LLAVES = {"transacciones": ["ref"], "valuacion_cetes": ["fecha", "instrumento_clave"],
          "precios_usados": ["fecha", "instrumento_clave"],
          "posiciones": ["nivel", "fecha_corte", "portafolio", "instrumento_clave", "intermediario"],
          "serie_diaria": ["fecha", "portafolio"], "indice_por_periodo": ["periodo", "fecha"],
          "indicadores": ["periodo"], "rechazos_rf28": ["id"], "casos_positivos": ["id"]}
# CSV fuera de esperados/ y secciones en las que solo se comparan las columnas que tiene la hoja
RUTA_CSV = {"rechazos_rf28": RAIZ / "rechazos_rf28.csv", "casos_positivos": RAIZ / "casos_positivos.csv"}
SOLO_COLUMNAS_DE_LA_HOJA = {"rechazos_rf28", "casos_positivos"}

TIPO = {}
for c in ("fecha vencimiento inf_fecha_colocacion sup_fecha_colocacion fecha_cierre fecha_corte fecha_precio "
          "fecha_colocacion_inf fecha_colocacion_sup inicio fin fecha_base fecha_pico fecha_valle").split():
    TIPO[c] = "fecha"
TIPO["fecha_saldo_negativo"] = "fecha"
for c in ("ref tipo instrumento_clave caso sin_cierre_ese_dia nivel portafolio intermediario tiene_transacciones "
          "periodo id resultado_esperado regla").split():
    TIPO[c] = "texto"
for c in ("d inf_plazo_nominal inf_plazo_efectivo sup_plazo_nominal sup_plazo_efectivo titulos dias saldo_resultante "
          "saldo_previo compras_del_dia ventas_del_dia saldo_si_ventas_antes_que_compras").split():
    TIPO[c] = "entero (títulos, días, plazos, saldos)"
for c in ("monto aportacion retiro costo_total valor resultado_no_realizado aportaciones retiros_ventas "
          "retiros_vencimientos retiros valor_base valor_fin costo_total_fin resultado_no_realizado_fin").split():
    TIPO[c] = "importe (pesos)"
for c in "costo_por_titulo ingreso_por_titulo costo_unitario precio".split():
    TIPO[c] = "precio o costo por título (pesos)"
TIPO["p"] = "P de CETES a 7 decimales"
TIPO["p_sin_redondear"] = "P de CETES sin redondear"
for c in "inf_tasa sup_tasa r_por_ciento_anual".split():
    TIPO[c] = "tasa r de RF-26b (% anual)"
TIPO["r"] = "r(t) (fracción)"
TIPO["r_por_ciento"] = "r(t) (%)"
for c in ("indice", "maximo"):
    TIPO[c] = "índice I(t) y máximo"
TIPO["caida"] = "caída I/máx − 1 (fracción)"
TIPO["twr"] = "TWR (fracción)"
TIPO["twr_por_ciento"] = "TWR (%)"
TIPO["maxima_caida"] = "máxima caída (fracción)"
TIPO["maxima_caida_por_ciento"] = "máxima caída (%)"
EXACTOS = {"fecha", "texto", "entero (títulos, días, plazos, saldos)", "P de CETES a 7 decimales"}
FRACCION = {"r(t) (fracción)", "índice I(t) y máximo", "caída I/máx − 1 (fracción)", "TWR (fracción)",
            "máxima caída (fracción)"}


def norm_libro(v):
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, bool):
        return "sí" if v else "no"
    return v


def lee_libro():
    wb = load_workbook(LIBRO, data_only=True, read_only=True)
    ws = wb["Esperados"]
    secciones, actual, columnas = {}, None, None
    for fila in ws.iter_rows(values_only=True):
        a = fila[0] if fila else None
        if isinstance(a, str) and a.startswith("## "):
            actual, columnas = a[3:].strip(), None
            secciones[actual] = []
            continue
        if actual is None:
            continue
        if columnas is None:
            columnas = [c for c in fila if c is not None]
            continue
        if all(v is None or v == "" for v in fila):
            actual = None
            continue
        secciones[actual].append({c: norm_libro(fila[i]) for i, c in enumerate(columnas)})
    return secciones


def lee_csv(nombre):
    with open(RUTA_CSV.get(nombre, ESP / f"{nombre}.csv"), encoding="utf-8") as fh:
        return [{k: (v if v != "" else None) for k, v in r.items()} for r in csv.DictReader(fh)]


def llave(fila, cols, origen):
    out = []
    for c in cols:
        v = fila[c]
        if v is None:
            out.append("")
        elif origen == "csv" or isinstance(v, str):
            out.append(str(v))
        elif isinstance(v, date):
            out.append(v.isoformat())
        else:
            out.append(str(v))
    return tuple(out)


libro = lee_libro()
res = {}           # tipo -> dict(n, max, donde, malos)
problemas = []


def anota(tipo, dif, donde, malo=False):
    r = res.setdefault(tipo, {"n": 0, "max": Decimal(0), "donde": "", "malos": 0, "digito": 0})
    r["n"] += 1
    if dif is not None and dif >= r["max"]:
        if dif > r["max"] or not r["donde"]:
            r["donde"] = donde
        r["max"] = dif
    if malo:
        r["malos"] += 1
        problemas.append(donde)


for seccion, cols in LLAVES.items():
    filas_csv = lee_csv(seccion)
    filas_xl = libro.get(seccion, [])
    idx_xl = {llave(r, cols, "xl"): r for r in filas_xl}
    if len(idx_xl) != len(filas_xl) or len(filas_xl) != len(filas_csv):
        problemas.append(f"{seccion}: {len(filas_csv)} renglones en CSV y {len(filas_xl)} en la hoja")
    for rc in filas_csv:
        k = llave(rc, cols, "csv")
        rx = idx_xl.get(k)
        if rx is None:
            problemas.append(f"{seccion}: falta en la hoja el renglón {k}")
            continue
        for c, vc in rc.items():
            if seccion in SOLO_COLUMNAS_DE_LA_HOJA and c not in rx:
                continue
            tipo = TIPO[c]
            vx = rx.get(c)
            donde = f"{seccion} {'/'.join(k)} [{c}]"
            if vc is None or vx is None:
                anota(tipo, None, donde, malo=not (vc is None and vx is None))
                continue
            if tipo == "fecha":
                vx_d = vx if isinstance(vx, date) else None
                anota(tipo, None, donde, malo=(vx_d is None or vx_d.isoformat() != vc))
            elif tipo == "texto":
                anota(tipo, None, donde, malo=(str(vx) != vc))
            elif tipo.startswith("entero"):
                anota(tipo, Decimal(abs(int(vc) - vx)), donde, malo=(int(vc) != vx))
            else:
                dif = abs(Decimal(vx) - Decimal(vc))      # Decimal(float) es exacto: diferencia sin redondear
                if tipo == "P de CETES a 7 decimales":
                    igual = format(Decimal(vx), ".7f") == vc
                    res.setdefault(tipo, {"n": 0, "max": Decimal(0), "donde": "", "malos": 0, "digito": 0})
                    res[tipo]["digito"] += 1 if igual else 0
                    anota(tipo, dif, donde, malo=not igual)
                else:
                    anota(tipo, dif, donde, malo=dif > UMBRAL)

orden = ["importe (pesos)", "precio o costo por título (pesos)", "P de CETES a 7 decimales",
         "P de CETES sin redondear", "tasa r de RF-26b (% anual)", "r(t) (fracción)", "r(t) (%)",
         "índice I(t) y máximo", "caída I/máx − 1 (fracción)", "TWR (fracción)", "TWR (%)",
         "máxima caída (fracción)", "máxima caída (%)", "entero (títulos, días, plazos, saldos)", "fecha", "texto"]
filas = []
print("Diferencia absoluta máxima entre la hoja (Excel) y Python (decimal), por tipo de cifra:")
for tipo in orden:
    r = res.get(tipo)
    if not r:
        continue
    if tipo in ("fecha", "texto") or tipo.startswith("entero"):
        txt_max = "coincidencia exacta" if r["malos"] == 0 else f"{r['malos']} distintas"
        pp = ""
        r["donde"] = ""
    else:
        txt_max = f"{r['max']:.3E}"
        pp = f"{r['max'] * 100:.3E}" if tipo in FRACCION else ""
    extra = f"; {r['digito']} de {r['n']} iguales dígito a dígito" if tipo == "P de CETES a 7 decimales" else ""
    print(f"  {tipo:<38} n = {r['n']:>4}  máx = {txt_max:<22}{('  (' + pp + ' pp)') if pp else ''}"
          f"{('  en ' + r['donde']) if r['donde'] and tipo not in ('fecha', 'texto') else ''}{extra}")
    filas.append([tipo, r["n"], txt_max, pp, r["donde"] if tipo not in ("fecha", "texto") else "",
                  "exacta" if tipo in EXACTOS else "1e-6", "ok" if r["malos"] == 0 else "FALLA"])
with open(RAIZ / "comparacion_vias.csv", "w", encoding="utf-8", newline="") as fh:
    w = csv.writer(fh, lineterminator="\n")
    w.writerow(["tipo_de_cifra", "cifras_comparadas", "diferencia_maxima", "diferencia_maxima_en_pp",
                "donde", "criterio", "estado"])
    w.writerows(filas)
if problemas:
    print("\nPROBLEMAS:")
    for p in problemas[:40]:
        print("  -", p)
    sys.exit(1)
print("\nSin diferencias fuera del criterio.")
